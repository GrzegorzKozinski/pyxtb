"""Reading XTB exports (spec §3, §6): zip of xlsx files -> operations per sub-account."""
from __future__ import annotations

import io
import re
import warnings
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime, timezone
from decimal import Decimal, InvalidOperation
from zoneinfo import ZoneInfo

import openpyxl

WARSAW = ZoneInfo("Europe/Warsaw")
ACCOUNT_ORDER = ("XTB", "EUR", "USD", "Z_IKE")

# file prefix -> (value of the Konto column, sub-account currency)
ACCOUNT_PREFIXES = {
    "PLN": ("XTB", "PLN"),
    "EUR": ("EUR", "EUR"),
    "USD": ("USD", "USD"),
    "IKE": ("Z_IKE", "PLN"),
}

# export `Type` -> how the builder treats it (spec §3.3)
KIND_BY_TYPE = {
    "stock purchase": "trade",
    "stock sell": "trade",
    "dividend": "dividend",
    "withholding tax": "dividend",
    "deposit": "deposit",
    "transfer": "transfer",
    "withdrawal": "withdrawal",
    "ike deposit": "ike_deposit",
    "subaccount transfer": "skip",
    "free funds interest": "skip",
    "free funds interest tax": "skip",
    "sec fee": "skip",
    "total": "total",
}


class ExportError(Exception):
    """Problem with a single file (the rest of the zip is still processed)."""


# eq=False: operations are compared and hashed by identity, so two identical
# export lines stay distinct and can be put in sets.
@dataclass(frozen=True, eq=False)
class Operation:
    account: str
    type: str
    instrument: str         # instrument name from the export
    ticker: str             # XTB symbol, e.g. TTWO.US
    time: datetime          # timezone-aware, UTC
    amount: Decimal
    comment: str

    @property
    def date(self) -> date:
        """Operation date in Europe/Warsaw (spec §4.3)."""
        return self.time.astimezone(WARSAW).date()

    @property
    def kind(self) -> str:
        return KIND_BY_TYPE.get(self.type.strip().lower(), "unknown")


@dataclass
class AccountData:
    account: str            # XTB / EUR / USD / Z_IKE
    currency: str           # sub-account currency
    source_file: str
    operations: list[Operation] = field(default_factory=list)
    total: Decimal | None = None   # the export's "Total" row (cash balance)
    warnings: list[str] = field(default_factory=list)

    @property
    def amount_sum(self) -> Decimal:
        return sum((o.amount for o in self.operations), Decimal(0))

    @property
    def balance_ok(self) -> bool | None:
        if self.total is None:
            return None
        cent = Decimal("0.01")
        return self.amount_sum.quantize(cent) == self.total.quantize(cent)


# ---------- value parsers ----------

def parse_amount(value) -> Decimal | None:
    """Float from xlsx or text ('1 234,56', '1,234.56', '-47,48')."""
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, Decimal):
        return value
    if isinstance(value, float):
        return Decimal(repr(value))
    if isinstance(value, int):
        return Decimal(value)
    s = str(value).strip().strip('"').replace("\u00a0", "").replace(" ", "")
    if not s:
        return None
    if "," in s and "." in s:   # the last separator is the decimal one
        s = s.replace(".", "").replace(",", ".") if s.rfind(",") > s.rfind(".") else s.replace(",", "")
    else:
        s = s.replace(",", ".")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


_TIME_FORMATS = (
    "%Y-%m-%d %H:%M:%S.%f", "%Y-%m-%d %H:%M:%S", "%Y-%m-%dT%H:%M:%S.%f", "%Y-%m-%dT%H:%M:%S",
    "%d.%m.%Y %H:%M:%S.%f", "%d.%m.%Y %H:%M:%S", "%d/%m/%Y %H:%M:%S", "%Y-%m-%d",
)


def parse_time(value) -> datetime | None:
    """Export times are UTC; naive datetimes are treated as UTC."""
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        dt = value
    elif isinstance(value, date):
        dt = datetime(value.year, value.month, value.day)
    else:
        dt = None
        for fmt in _TIME_FORMATS:
            try:
                dt = datetime.strptime(str(value).strip(), fmt)
                break
            except ValueError:
                continue
        if dt is None:
            return None
    return dt.replace(tzinfo=timezone.utc) if dt.tzinfo is None else dt.astimezone(timezone.utc)


def _text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


# ---------- trade comment (spec §6) ----------

_TRADE_RE = re.compile(
    r"^\s*(?P<action>OPEN|CLOSE)\s+(?P<side>BUY|SELL)\s+"
    r"(?P<qty>\d+(?:\.\d+)?)(?:\s*/\s*\d+(?:\.\d+)?)?"      # qty of this line [/ qty of the whole order]
    r"\s*@\s*(?P<price>\d+(?:\.\d+)?)",
    re.IGNORECASE,
)


@dataclass(frozen=True)
class TradeComment:
    action: str            # OPEN | CLOSE
    side: str              # BUY | SELL
    qty: Decimal           # quantity of this line (partial fill)
    price: Decimal

    @property
    def supported(self) -> bool:
        """OPEN BUY (purchase) and CLOSE BUY (sale); short positions are not supported."""
        return self.side == "BUY"


def parse_trade_comment(comment: str | None) -> TradeComment | None:
    m = _TRADE_RE.match(comment or "")
    if not m:
        return None
    return TradeComment(m["action"].upper(), m["side"].upper(), Decimal(m["qty"]), Decimal(m["price"]))


# ---------- files ----------

def account_from_filename(name: str) -> tuple[str, str] | None:
    """(account, currency) from the file name prefix, or None. B1 in the file is unreliable (spec §3.1)."""
    base = name.replace("\\", "/").rsplit("/", 1)[-1]
    return ACCOUNT_PREFIXES.get(base.split("_", 1)[0].upper())


def read_cash_operations(data: bytes, account: str, currency: str, source_file: str) -> AccountData:
    """Reads the Cash Operations sheet of one xlsx file."""
    try:
        # Not read_only: XTB exports have a wrong <dimension> tag, so read_only sees only part of the sheet.
        with warnings.catch_warnings():
            warnings.simplefilter("ignore", UserWarning)    # "Workbook contains no default style"
            wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    except Exception as e:  # corrupted / not an xlsx
        raise ExportError(f"nie można otworzyć pliku xlsx: {e}") from e

    sheet = next((n for n in wb.sheetnames if n.strip().lower() == "cash operations"), None)
    if sheet is None:
        raise ExportError(f"brak zakładki „Cash Operations” (są: {', '.join(wb.sheetnames)})")
    rows = list(wb[sheet].iter_rows(values_only=True))
    wb.close()

    # the spec says headers are in row 5; look for them in the first 10 rows
    header_idx = next((i for i, row in enumerate(rows[:10])
                       if {"type", "amount", "time"} <= {str(c).strip().lower() for c in row if c is not None}), None)
    if header_idx is None:
        raise ExportError(f"nie znaleziono wiersza nagłówków (Type/Amount/Time) w Cash Operations (wierszy: {len(rows)})")
    cols = {str(c).strip().lower(): j for j, c in enumerate(rows[header_idx]) if c is not None}

    def cell(row, name):
        j = cols.get(name)
        return row[j] if j is not None and j < len(row) else None

    result = AccountData(account=account, currency=currency, source_file=source_file)
    for n, row in enumerate(rows[header_idx + 1:], start=header_idx + 2):
        if row is None or all(c is None or str(c).strip() == "" for c in row):
            continue
        op_type = _text(cell(row, "type"))
        amount = parse_amount(cell(row, "amount"))
        if op_type.lower() == "total":
            result.total = amount
            continue
        time = parse_time(cell(row, "time"))
        if amount is None or time is None:
            result.warnings.append(f"{source_file}, wiersz {n}: nieczytelny Amount/Time – pominięto ({op_type or '?'})")
            continue
        result.operations.append(Operation(
            account=account, type=op_type, instrument=_text(cell(row, "instrument")),
            ticker=_text(cell(row, "ticker")), time=time, amount=amount, comment=_text(cell(row, "comment"))))
    return result


def read_zip(zip_path: str) -> tuple[dict[str, AccountData], list[str]]:
    """Reads the zip. Returns ({account: AccountData}, file-level problems)."""
    accounts: dict[str, AccountData] = {}
    problems: list[str] = []
    with zipfile.ZipFile(zip_path) as zf:
        for name in sorted(zf.namelist()):
            base = name.replace("\\", "/").rsplit("/", 1)[-1]
            if not base or base.startswith("~$") or not base.lower().endswith(".xlsx"):
                continue
            acc = account_from_filename(base)
            if acc is None:
                problems.append(f"{base}: nieznany prefiks pliku – pominięto")
                continue
            account, currency = acc
            if account in accounts:
                problems.append(f"{base}: konto {account} już wczytane z {accounts[account].source_file} – pominięto")
                continue
            try:
                accounts[account] = read_cash_operations(zf.read(name), account, currency, base)
            except ExportError as e:
                problems.append(f"{base}: {e}")
            else:
                problems.extend(accounts[account].warnings)    # unreadable rows etc. end up in the report
    return accounts, problems
